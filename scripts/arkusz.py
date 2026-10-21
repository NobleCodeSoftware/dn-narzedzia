#!/usr/bin/env python3
"""Generuje darmowy arkusz „Ewidencja sprzedaży i limit DN” (XLSX) do public/pobierz/.

Arkusz działa w Excelu, LibreOffice Calc i Arkuszach Google: same formuły, bez makr,
bez połączeń z siecią, bez łączy zewnętrznych i bez funkcji HYPERLINK.
Wymaga openpyxl (licencja MIT): pip install openpyxl

Użycie (z katalogu głównego repozytorium):
    python3 scripts/arkusz.py            # zapisuje public/pobierz/arkusz-limit-dn.xlsx
    python3 scripts/arkusz.py --sprawdz  # zapisuje i sprawdza plik

Plik jest deterministyczny (stałe daty w metadanych i w archiwum ZIP), więc ponowne
uruchomienie bez zmian w skrypcie daje identyczny plik.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

KATALOG = Path(__file__).resolve().parent.parent
WYJSCIE = KATALOG / "public" / "pobierz" / "arkusz-limit-dn.xlsx"
STALA_DATA = dt.datetime(2026, 10, 21, 0, 0, 0)

# Minimalne wynagrodzenie (zł) i źródło. Limit kwartalny = 225% tej kwoty
# (art. 5 ust. 1 ustawy z 6 marca 2018 r. – Prawo przedsiębiorców). Te same kwoty są w public/js/logic.js.
PLACA = [
    (2026, 4806, "Dz.U. 2025 poz. 1242 (rozporządzenie Rady Ministrów z 11.09.2025)"),
    (2027, 4950, "Dz.U. 2026 poz. 1213 (rozporządzenie Rady Ministrów z 14.09.2026)"),
]
LATA_DO_UZUPELNIENIA = range(2028, 2036)
LATA_PODSUMOWANIA = (2026, 2027, 2028)

PIERWSZY, OSTATNI = 5, 1004  # wiersze na sprzedaż w zakładce „Sprzedaż”
LIMITY_ZAKRES = "Limity!$A$5:$C$14"

STOPKA = "Eksperyment NobleCode, prowadzony z udziałem AI · Nie stanowi doradztwa podatkowego · kontakt@noblecode.pl"

# Kolory (tło, tekst) dla progów: ten sam układ co na stronie.
PROGI = [
    (1.00, "B42318", "FFFFFF"),  # 100% i więcej
    (0.95, "FDE2DF", "8A1C12"),
    (0.85, "FFE3C2", "7A3E00"),
    (0.70, "FFF3C4", "6B5200"),
]

NAGLOWEK_TLO = PatternFill("solid", fgColor="1F5F8B")
NAGLOWEK_FONT = Font(bold=True, color="FFFFFF")
LICZONE_TLO = PatternFill("solid", fgColor="F1F0EC")
CIENKA = Side(style="thin", color="DCDAD5")
RAMKA = Border(left=CIENKA, right=CIENKA, top=CIENKA, bottom=CIENKA)
ZAWIJAJ = Alignment(wrap_text=True, vertical="top")

FORMAT_DATY = "yyyy-mm-dd"
FORMAT_KWOTY = "#,##0.00"
FORMAT_ZL = '#,##0.00 "zł"'
FORMAT_PROC = "0.0%"


def drabinka_ostrzezen(ref: str) -> str:
    """Tekst ostrzeżenia dla komórki z ułamkiem limitu (ref jest liczbą)."""
    return (
        f'IF({ref}>1,"Limit przekroczony",IF({ref}>=1,"Limit wykorzystany (100%)",'
        f'IF({ref}>=0.95,"Co najmniej 95% limitu",IF({ref}>=0.85,"Co najmniej 85% limitu",'
        f'IF({ref}>=0.7,"Co najmniej 70% limitu","W limicie")))))'
    )


def kolory_progow(ws, zakres: str, kolumna_proc: str, pierwszy_wiersz: int) -> None:
    ref = f"${kolumna_proc}{pierwszy_wiersz}"
    for prog, tlo, tekst in PROGI:
        ws.conditional_formatting.add(
            zakres,
            FormulaRule(
                formula=[f"AND(ISNUMBER({ref}),{ref}>={prog})"],
                fill=PatternFill("solid", fgColor=tlo, bgColor=tlo),
                font=Font(color=tekst, bold=True),
                stopIfTrue=True,
            ),
        )


def druk_poziomo(ws) -> None:
    """Wydruk: poziomo, wszystkie kolumny na szerokość strony."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"


def naglowki(ws, wiersz: int, nazwy: list[str]) -> None:
    for i, nazwa in enumerate(nazwy, start=1):
        c = ws.cell(row=wiersz, column=i, value=nazwa)
        c.fill = NAGLOWEK_TLO
        c.font = NAGLOWEK_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = RAMKA
    ws.row_dimensions[wiersz].height = 32


def zakladka_sprzedaz(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Sprzedaż"
    ws["A1"] = "Ewidencja sprzedaży i limit działalności nierejestrowanej"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Jeden wiersz = jedna sprzedaż. Wypełniasz kolumny A–E, szare kolumny liczą się same. "
                "Zwrot wpisz jako kwotę ujemną. Jak znaleźć kurs NBP: zakładka „Instrukcja”.")
    ws.merge_cells("A2:I2")
    ws["A2"].alignment = ZAWIJAJ
    ws.row_dimensions[2].height = 34
    ws["A3"] = STOPKA
    ws["A3"].font = Font(size=9, color="5B616B")
    druk_poziomo(ws)
    naglowki(ws, 4, ["Data", "Opis", "Kwota", "Waluta", "Kurs NBP", "Kwota w PLN",
                     "Suma kwartału (narastająco)", "% limitu", "Ostrzeżenie"])
    zakres_f = f"$F${PIERWSZY}:$F${OSTATNI}"
    zakres_a = f"$A${PIERWSZY}:$A${OSTATNI}"
    for r in range(PIERWSZY, OSTATNI + 1):
        ws[f"A{r}"].number_format = FORMAT_DATY
        ws[f"C{r}"].number_format = FORMAT_KWOTY
        ws[f"E{r}"].number_format = "0.0000"
        ws[f"F{r}"] = (f'=IF(C{r}="","",IF(OR(D{r}="",D{r}="PLN"),ROUND(C{r},2),'
                       f'IF(E{r}="","brak kursu",ROUND(C{r}*E{r},2))))')
        ws[f"G{r}"] = (f'=IF(OR(A{r}="",NOT(ISNUMBER(F{r}))),"",IF(NOT(ISNUMBER(A{r})),"sprawdź datę",'
                       f'SUMIFS({zakres_f},{zakres_a},">="&DATE(YEAR(A{r}),INT((MONTH(A{r})-1)/3)*3+1,1),'
                       f'{zakres_a},"<="&A{r})))')
        ws[f"H{r}"] = (f'=IF(NOT(ISNUMBER(G{r})),"",'
                       f'IFERROR(G{r}/VLOOKUP(YEAR(A{r}),{LIMITY_ZAKRES},3,FALSE),"brak limitu"))')
        ws[f"I{r}"] = (f'=IF(ISNUMBER(H{r}),{drabinka_ostrzezen(f"H{r}")},'
                       f'IF(F{r}="brak kursu","Wpisz kurs NBP w kolumnie E",'
                       f'IF(G{r}="sprawdź datę","Wpisz datę jako dzień, np. 2026-10-21",'
                       f'IF(H{r}="brak limitu","Wpisz płacę minimalną na ten rok w zakładce Limity",""))))')
        ws[f"F{r}"].number_format = FORMAT_ZL
        ws[f"G{r}"].number_format = FORMAT_ZL
        ws[f"H{r}"].number_format = FORMAT_PROC
        for kol in "FGHI":
            ws[f"{kol}{r}"].fill = LICZONE_TLO
    kolory_progow(ws, f"H{PIERWSZY}:I{OSTATNI}", "H", PIERWSZY)
    ws.conditional_formatting.add(
        f"F{PIERWSZY}:F{OSTATNI}",
        FormulaRule(formula=[f'$F{PIERWSZY}="brak kursu"'],
                    fill=PatternFill("solid", fgColor="FFE3C2", bgColor="FFE3C2"),
                    font=Font(color="7A3E00", bold=True)),
    )
    for kol, szer in zip("ABCDEFGHI", (12, 34, 12, 9, 11, 14, 16, 10, 30)):
        ws.column_dimensions[kol].width = szer
    ws.freeze_panes = f"A{PIERWSZY}"

    dv_data = DataValidation(type="date", operator="between", formula1="DATE(2026,1,1)", formula2="DATE(2035,12,31)",
                             errorStyle="warning", showErrorMessage=True, allow_blank=True,
                             errorTitle="Data", error="Wpisz datę sprzedaży, np. 2026-10-21.")
    dv_kwota = DataValidation(type="decimal", operator="between", formula1="-1000000", formula2="1000000",
                              errorStyle="warning", showErrorMessage=True, allow_blank=True,
                              errorTitle="Kwota", error="Wpisz kwotę liczbą, np. 49,90. Zwrot jako kwotę ujemną.")
    dv_waluta = DataValidation(type="list", formula1='"PLN,EUR,USD,GBP,CHF,CZK,SEK,NOK,DKK,CAD,AUD"',
                               errorStyle="information", showErrorMessage=True, allow_blank=True,
                               errorTitle="Waluta", error="Inna waluta? Wpisz jej trzyliterowy kod, np. JPY.")
    dv_kurs = DataValidation(type="decimal", operator="greaterThan", formula1="0",
                             errorStyle="warning", showErrorMessage=True, allow_blank=True,
                             errorTitle="Kurs NBP", error="Kurs to liczba większa od zera, np. 4,3641.")
    for dv, kol in ((dv_data, "A"), (dv_kwota, "C"), (dv_waluta, "D"), (dv_kurs, "E")):
        ws.add_data_validation(dv)
        dv.add(f"{kol}{PIERWSZY}:{kol}{OSTATNI}")


def zakladka_podsumowanie(wb: Workbook) -> None:
    ws = wb.create_sheet("Podsumowanie")
    ws["A1"] = "Podsumowanie kwartałów"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Liczy się samo z zakładki „Sprzedaż”. Kolory: 70%, 85%, 95% i 100% limitu."
    druk_poziomo(ws)
    naglowki(ws, 4, ["Rok", "Kwartał", "Od", "Do", "Przychód", "Limit na kwartał", "% limitu",
                     "Zostało do limitu", "Ostrzeżenie"])
    s = "'Sprzedaż'!"
    r = 5
    for rok in LATA_PODSUMOWANIA:
        for q in range(1, 5):
            ws[f"A{r}"] = rok
            ws[f"B{r}"] = ["I", "II", "III", "IV"][q - 1]
            ws[f"C{r}"] = f"=DATE(A{r},{(q - 1) * 3 + 1},1)"
            ws[f"D{r}"] = f"=EOMONTH(C{r},2)"
            ws[f"E{r}"] = (f'=SUMIFS({s}$F${PIERWSZY}:$F${OSTATNI},{s}$A${PIERWSZY}:$A${OSTATNI},">="&C{r},'
                           f'{s}$A${PIERWSZY}:$A${OSTATNI},"<="&D{r})')
            ws[f"F{r}"] = f'=IFERROR(VLOOKUP(A{r},{LIMITY_ZAKRES},3,FALSE),"")'
            ws[f"G{r}"] = f'=IF(ISNUMBER(F{r}),E{r}/F{r},"")'
            ws[f"H{r}"] = f'=IF(ISNUMBER(F{r}),F{r}-E{r},"")'
            bez_kursu = (f'COUNTIFS({s}$F${PIERWSZY}:$F${OSTATNI},"brak kursu",{s}$A${PIERWSZY}:$A${OSTATNI},">="&C{r},'
                         f'{s}$A${PIERWSZY}:$A${OSTATNI},"<="&D{r})')
            ws[f"I{r}"] = (f'=IF(ISNUMBER(G{r}),{drabinka_ostrzezen(f"G{r}")},'
                           f'"Wpisz płacę minimalną na ten rok w zakładce Limity")'
                           f'&IF({bez_kursu}>0," · wierszy bez kursu NBP: "&{bez_kursu},"")')
            for kol, fmt in (("C", FORMAT_DATY), ("D", FORMAT_DATY), ("E", FORMAT_ZL), ("F", FORMAT_ZL),
                             ("G", FORMAT_PROC), ("H", FORMAT_ZL)):
                ws[f"{kol}{r}"].number_format = fmt
            for kol in "ABCDEFGHI":
                ws[f"{kol}{r}"].border = RAMKA
            r += 1
    kolory_progow(ws, f"G5:I{r - 1}", "G", 5)
    for kol, szer in zip("ABCDEFGHI", (7, 9, 12, 12, 14, 16, 10, 17, 30)):
        ws.column_dimensions[kol].width = szer
    ws.freeze_panes = "A5"


def zakladka_limity(wb: Workbook) -> None:
    ws = wb.create_sheet("Limity")
    ws["A1"] = "Limit na kwartał = 225% minimalnego wynagrodzenia"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Podstawa: art. 5 ust. 1 ustawy z 6 marca 2018 r. – Prawo przedsiębiorców. "
                "Kwotę na kolejny rok wpisz w kolumnie B, gdy rząd ogłosi minimalne wynagrodzenie (zwykle we wrześniu).")
    ws["A2"].alignment = Alignment(wrap_text=False)
    naglowki(ws, 4, ["Rok", "Minimalne wynagrodzenie (zł)", "Limit na kwartał (zł)", "Źródło"])
    r = 5
    for rok, placa, zrodlo in PLACA:
        ws[f"A{r}"], ws[f"B{r}"], ws[f"D{r}"] = rok, placa, zrodlo
        r += 1
    for rok in LATA_DO_UZUPELNIENIA:
        ws[f"A{r}"], ws[f"D{r}"] = rok, "Wpisz po ogłoszeniu minimalnego wynagrodzenia"
        ws[f"D{r}"].font = Font(color="5B616B", italic=True)
        r += 1
    for w in range(5, r):
        ws[f"C{w}"] = f'=IF(B{w}="","",ROUND(B{w}*2.25,2))'
        ws[f"B{w}"].number_format = FORMAT_KWOTY
        ws[f"C{w}"].number_format = FORMAT_ZL
        ws[f"C{w}"].fill = LICZONE_TLO
        for kol in "ABCD":
            ws[f"{kol}{w}"].border = RAMKA
    for kol, szer in zip("ABCD", (7, 16, 16, 62)):
        ws.column_dimensions[kol].width = szer


INSTRUKCJA = [
    ("Jak korzystać z arkusza", "h"),
    ("1. W zakładce „Sprzedaż” wpisz każdą sprzedaż w osobnym wierszu: datę, krótki opis, kwotę i walutę.", ""),
    ("2. Sprzedaż w złotych: zostaw kolumnę „Kurs NBP” pustą. Sprzedaż w innej walucie: wpisz kurs NBP (niżej, jak go znaleźć).", ""),
    ("3. Szare kolumny liczą się same: kwota w PLN, suma kwartału narastająco, procent limitu i ostrzeżenie.", ""),
    ("4. Zwrot pieniędzy klientowi wpisz jako kwotę ujemną z datą zwrotu.", ""),
    ("5. Zakładka „Podsumowanie” pokazuje każdy kwartał: przychód, limit, procent i ile zostało.", ""),
    ("", ""),
    ("Co znaczą kolory", "h"),
    ("Żółty: co najmniej 70% limitu. Pomarańczowy: co najmniej 85%. Jasnoczerwony: co najmniej 95%.", ""),
    ("Czerwony: limit wykorzystany w 100% albo przekroczony.", ""),
    ("", ""),
    ("Jak znaleźć kurs średni NBP", "h"),
    ("Przychód w walucie obcej przelicza się na złote po kursie średnim NBP z ostatniego dnia roboczego przed dniem uzyskania przychodu (art. 11a ust. 1 ustawy o PIT).", ""),
    ("1. Ustal dzień roboczy przed dniem sprzedaży. Przykład: sprzedaż w poniedziałek, więc kurs z piątku. Jeśli tego dnia było święto, cofnij się do poprzedniego dnia roboczego.", ""),
    ("2. Otwórz stronę NBP: nbp.pl, potem „Statystyka i sprawozdawczość”, „Kursy”, „Archiwum kursów średnich – tabela A”. Wybierz ustaloną datę.", ""),
    ("   Kursy z bieżącego dnia są w „Tabela A kursów średnich walut obcych”: nbp.pl/statystyka-i-sprawozdawczosc/kursy/tabela-a/", ""),
    ("3. Znajdź swoją walutę i przepisz kurs średni do kolumny „Kurs NBP”, np. 4,3641.", ""),
    ("4. Uwaga na waluty podawane za 100 jednostek (np. 100 JPY, 100 HUF): podziel kurs przez 100 przed wpisaniem.", ""),
    ("Przykład (liczby przykładowe): sprzedaż 10 EUR, kurs 4,3641, kwota w PLN = 43,64 zł.", ""),
    ("", ""),
    ("O arkuszu", "h"),
    ("Arkusz nie ma makr i nie łączy się z internetem. Twoje dane zostają w pliku, który sam przechowujesz.", ""),
    ("Działa w Excelu, LibreOffice Calc i Arkuszach Google (Plik, Importuj).", ""),
    ("Kwoty limitu i źródła są w zakładce „Limity”. Sprawdź je w źródłach, zanim podejmiesz decyzję.", ""),
    ("Arkusz pomaga liczyć, ale nie zastępuje ewidencji wymaganej przepisami ani porady podatkowej.", ""),
    ("", ""),
    (STOPKA, "s"),
]


def zakladka_instrukcja(wb: Workbook) -> None:
    ws = wb.create_sheet("Instrukcja")
    ws["A1"] = "Instrukcja"
    ws["A1"].font = Font(bold=True, size=14)
    for i, (tekst, styl) in enumerate(INSTRUKCJA, start=3):
        c = ws.cell(row=i, column=1, value=tekst or None)
        c.alignment = ZAWIJAJ
        if styl == "h":
            c.font = Font(bold=True, size=12)
        elif styl == "s":
            c.font = Font(size=9, color="5B616B")
    ws.column_dimensions["A"].width = 110


def zbuduj() -> bytes:
    wb = Workbook()
    zakladka_sprzedaz(wb)
    zakladka_podsumowanie(wb)
    zakladka_limity(wb)
    zakladka_instrukcja(wb)
    p = wb.properties
    p.creator = p.lastModifiedBy = "NobleCode"
    p.title = "Ewidencja sprzedaży i limit działalności nierejestrowanej"
    p.created = p.modified = STALA_DATA
    bufor = io.BytesIO()
    wb.save(bufor)
    return deterministyczny_zip(bufor.getvalue())


def deterministyczny_zip(dane: bytes) -> bytes:
    """Przepakowuje ZIP ze stałą datą wpisów, żeby plik był powtarzalny."""
    wej = zipfile.ZipFile(io.BytesIO(dane))
    stempel = STALA_DATA.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    wyj_bufor = io.BytesIO()
    with zipfile.ZipFile(wyj_bufor, "w", zipfile.ZIP_DEFLATED) as wyj:
        for info in wej.infolist():
            nowy = zipfile.ZipInfo(info.filename, date_time=STALA_DATA.timetuple()[:6])
            nowy.compress_type = zipfile.ZIP_DEFLATED
            nowy.external_attr = 0o644 << 16
            tresc = wej.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl przy zapisie nadpisuje dcterms:modified bieżącym czasem.
                tresc = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", rb"\g<1>" + stempel + rb"\g<2>", tresc)
            wyj.writestr(nowy, tresc)
    return wyj_bufor.getvalue()


def sprawdz(sciezka: Path) -> None:
    """Kontrola pliku: otwiera się, brak makr, łączy zewnętrznych, HYPERLINK i adresów w formułach."""
    with zipfile.ZipFile(sciezka) as z:
        nazwy = z.namelist()
        assert not any("vbaProject" in n or "externalLink" in n or n.endswith(".bin") for n in nazwy), nazwy
        for n in nazwy:
            if n.endswith(".rels"):
                assert b'TargetMode="External"' not in z.read(n), f"łącze zewnętrzne w {n}"
    wb = load_workbook(sciezka)
    assert wb.sheetnames == ["Sprzedaż", "Podsumowanie", "Limity", "Instrukcja"], wb.sheetnames
    formuly = 0
    for ws in wb.worksheets:
        assert not ws._hyperlinks, f"hiperłącza w {ws.title}"
        for wiersz in ws.iter_rows():
            for c in wiersz:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formuly += 1
                    gorne = c.value.upper()
                    assert "HYPERLINK" not in gorne and "HTTP" not in gorne and "[" not in gorne, c.coordinate
                    assert "WEBSERVICE" not in gorne and "IMPORT" not in gorne, c.coordinate
    assert wb["Limity"]["B5"].value == 4806 and wb["Limity"]["B6"].value == 4950
    print(f"OK: {sciezka.relative_to(KATALOG)} ({sciezka.stat().st_size} B, {formuly} formuł, bez makr i łączy)")


def main() -> None:
    WYJSCIE.parent.mkdir(parents=True, exist_ok=True)
    WYJSCIE.write_bytes(zbuduj())
    print(f"Zapisano {WYJSCIE.relative_to(KATALOG)}")
    if "--sprawdz" in sys.argv:
        sprawdz(WYJSCIE)


if __name__ == "__main__":
    main()
